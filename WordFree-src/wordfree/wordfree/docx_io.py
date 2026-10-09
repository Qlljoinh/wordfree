"""Открытие и сохранение документов .docx (через python-docx)."""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

from .docwalk import iter_blocks, parse_char_tags, px_to_pt

ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
ALIGN_BACK = {v: k for k, v in ALIGN_MAP.items()}

STYLE_MAP = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}
STYLE_BACK = {v: k for k, v in STYLE_MAP.items()}

HL_MAP = {
    "#ffff00": "YELLOW",
    "#00ff00": "BRIGHT_GREEN",
    "#00ffff": "TURQUOZE",
    "#ff00ff": "PINK",
    "#ffc000": "DARK_YELLOW",
    "#0000ff": "BLUE",
    "#ff0000": "RED",
    "#000000": "BLACK",
    "#ffffff": "WHITE",
    "#00ff80": "BRIGHT_GREEN",
    "#ff8080": "PINK",
}

_LEAD_BULLET_RE = re.compile(r"^\s*•\s*")
_LEAD_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s*")


def _tag_name_for_font(family):
    return "font_" + family.replace(" ", "_")


def _apply_run_format(run, fmt):
    if fmt.get("bold"):
        run.bold = True
    if fmt.get("italic"):
        run.italic = True
    if fmt.get("underline"):
        run.underline = True
    if fmt.get("strike"):
        run.font.strike = True
    if fmt.get("font"):
        run.font.name = fmt["font"]
    if fmt.get("size"):
        run.font.size = Pt(fmt["size"])
    if fmt.get("color"):
        try:
            run.font.color.rgb = RGBColor.from_string(fmt["color"].lstrip("#"))
        except Exception:
            pass
    if fmt.get("highlight"):
        name = HL_MAP.get(fmt["highlight"].lower())
        if name:
            try:
                run.font.highlight_color = getattr(WD_COLOR_INDEX, name)
            except Exception:
                pass


def save_docx(text, tables, images, path, title=""):
    """Сохранить содержимое Text-виджета в .docx.

    text   — tk.Text виджет редактора
    tables — реестр таблиц {id: {'rows','cols','cells': [[tk.Entry...]]}}
    images — реестр картинок {name: {'path','width_in'}}
    """
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def add_paragraph_block(block):
        p = doc.add_paragraph()
        pf = block["pfmt"]
        p.alignment = ALIGN_MAP.get(pf["align"], WD_ALIGN_PARAGRAPH.LEFT)
        if pf["style"] in STYLE_MAP:
            p.style = doc.styles[STYLE_MAP[pf["style"]]]
        elif pf["list"] == "bullet":
            p.style = doc.styles["List Bullet"]
        elif pf["list"] == "number":
            p.style = doc.styles["List Number"]
        p.paragraph_format.space_after = Pt(pf.get("space_after") or 0)
        if pf.get("line_spacing"):
            p.paragraph_format.line_spacing = pf["line_spacing"]
        if pf.get("indent_left"):
            p.paragraph_format.left_indent = Pt(px_to_pt(pf["indent_left"]))
        if pf.get("indent_first"):
            p.paragraph_format.first_line_indent = Pt(px_to_pt(pf["indent_first"]))
        if pf.get("page_break"):
            p.paragraph_format.page_break_before = True
        for item in block["flow"]:
            if item["kind"] == "text":
                txt = item["text"]
                if pf["list"] == "bullet":
                    txt = _LEAD_BULLET_RE.sub("", txt, count=1)
                elif pf["list"] == "number":
                    txt = _LEAD_NUMBER_RE.sub("", txt, count=1)
                if txt == "":
                    continue
                run = p.add_run(txt)
                _apply_run_format(run, item["fmt"])
            elif item["kind"] == "image":
                info = images.get(item["name"])
                if not info or not os.path.exists(info["path"]):
                    continue
                run = p.add_run()
                try:
                    run.add_picture(info["path"], width=Inches(info.get("width_in") or 4.0))
                except Exception as e:
                    print(f"[docx] картинка не вставлена: {e}")
        return p

    has_content = False
    for block in iter_blocks(text):
        has_content = True
        if block["type"] == "table":
            tinfo = tables.get(block["table_id"])
            if not tinfo:
                continue
            rows, cols = tinfo["rows"], tinfo["cols"]
            tbl = doc.add_table(rows=rows, cols=cols)
            tbl.style = "Table Grid"
            for i in range(rows):
                for j in range(cols):
                    try:
                        cell_text = tinfo["cells"][i][j].get()
                    except Exception:
                        cell_text = ""
                    tbl.cell(i, j).text = cell_text
            doc.add_paragraph()  # разделитель после таблицы
        else:
            add_paragraph_block(block)

    if not has_content:
        doc.add_paragraph()

    try:
        doc.core_properties.title = title or os.path.splitext(os.path.basename(path))[0]
        doc.core_properties.author = "WordFree"
    except Exception:
        pass
    doc.save(path)


def open_docx(text, tables, images, path, insert_image_cb, create_table_cb, temp_dir):
    """Загрузить .docx в Text-виджет редактора.

    insert_image_cb(path, width_in) — вставить картинку в конец документа
    create_table_cb(rows, cols, cell_texts) — создать таблицу в конце документа
    """
    doc = Document(path)
    # очистка содержимого
    text.delete("1.0", "end")

    def ensure_newline():
        # гарантировать, что следующий абзац начнётся с новой строки
        if not text.index("end-1c").endswith(".0"):
            text.insert("end", "\n")

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            ensure_newline()
            line0 = int(text.index("end-1c").split(".")[0])
            _load_paragraph(text, p, doc, images, insert_image_cb, temp_dir)
            # абзацные теги на все строки абзаца (обычно одна)
            pf = p.paragraph_format
            style_name = p.style.name if p.style else "Normal"
            tags = []
            if style_name in STYLE_BACK:
                tags.append("style_" + STYLE_BACK[style_name])
            elif style_name == "List Bullet":
                tags.append("list_bullet")
            elif style_name == "List Number":
                tags.append("list_number")
            if p.alignment in ALIGN_BACK:
                tags.append("al_" + ALIGN_BACK[p.alignment])
            if pf.line_spacing:
                tags.append(f"sp_line_{pf.line_spacing}")
            if pf.space_after is not None and pf.space_after.pt:
                tags.append(f"sp_after_{int(pf.space_after.pt)}")
            if pf.left_indent is not None and pf.left_indent.pt:
                tags.append(f"ind_left_{int(pf.left_indent.pt * 96 / 72)}")
            if pf.first_line_indent is not None and pf.first_line_indent.pt:
                tags.append(f"ind_first_{int(pf.first_line_indent.pt * 96 / 72)}")
            if pf.page_break_before:
                tags.append("pb")
            last_line = int(text.index("end-1c").split(".")[0])
            for ln in range(line0, last_line + 1):
                for t in tags:
                    text.tag_add(t, f"{ln}.0", f"{ln}.end")
            text.insert("end", "\n")
        elif child.tag == qn("w:tbl"):
            tbl = DocxTable(child, doc)
            rows = len(tbl.rows)
            cols = len(tbl.columns)
            cell_texts = []
            for row in tbl.rows:
                r = []
                for cell in row.cells:
                    r.append("\n".join(pp.text for pp in cell.paragraphs))
                cell_texts.append(r)
            ensure_newline()
            create_table_cb(rows, cols, cell_texts)
            text.insert("end", "\n")


def _load_paragraph(text, p, doc, images, insert_image_cb, temp_dir):
    """Вставить один абзац (с run-форматированием и картинками) в конец виджета."""
    # картинки абзаца
    blips = list(p._p.iter(qn("a:blip")))
    extents = list(p._p.iter(qn("wp:extent")))
    image_rids = []
    for idx, blip in enumerate(blips):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        try:
            rel = doc.part.related_parts.get(rid)
            img = rel.image
        except Exception:
            continue
        try:
            ext = img.ext or "png"
            fname = os.path.join(temp_dir, f"img_{len(images)}_{idx}.{ext}")
            with open(fname, "wb") as f:
                f.write(img.blob)
            width_in = 4.0
            if idx < len(extents):
                cx = int(extents[idx].get("cx", 0))
                if cx:
                    width_in = cx / 914400.0
            insert_image_cb(fname, width_in)
            image_rids.append(rid)
        except Exception as e:
            print(f"[docx] картинка не загружена: {e}")

    for run in p.runs:
        fmt = {}
        f = run.font
        if f.bold:
            fmt["bold"] = True
        if f.italic:
            fmt["italic"] = True
        if f.underline:
            fmt["underline"] = True
        if f.strike:
            fmt["strike"] = True
        if f.name:
            fmt["font"] = f.name
        if f.size is not None:
            fmt["size"] = int(f.size.pt)
        if f.color is not None and f.color.rgb is not None:
            fmt["color"] = "#" + str(f.color.rgb)
        if f.highlight_color is not None and str(f.highlight_color) != "WD_COLOR_INDEX.AUTO":
            try:
                fmt["highlight"] = "#" + str(f.highlight_color.value)[-6:] if hasattr(f.highlight_color, "value") else None
            except Exception:
                fmt["highlight"] = None
        txt = run.text.replace("\t", "    ").replace("\n", "").replace("\r", "")
        if txt:
            tags = _tags_from_fmt(fmt)
            text.insert("end", txt, tuple(tags))


def _tags_from_fmt(fmt):
    tags = []
    if fmt.get("bold"):
        tags.append("b")
    if fmt.get("italic"):
        tags.append("i")
    if fmt.get("underline"):
        tags.append("u")
    if fmt.get("strike"):
        tags.append("s")
    if fmt.get("font"):
        tags.append(_tag_name_for_font(fmt["font"]))
    if fmt.get("size"):
        tags.append(f"size_{fmt['size']}")
    if fmt.get("color"):
        tags.append("color_" + fmt["color"].lstrip("#").lower())
    if fmt.get("highlight"):
        tags.append("hl_" + fmt["highlight"].lstrip("#").lower())
    return tags
