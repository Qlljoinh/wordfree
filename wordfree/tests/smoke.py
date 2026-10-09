"""Смоук-тест WordFree: запуск UI под Xvfb и проверка ключевых функций."""
import os
import sys
import tempfile
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from wordfree.app import WordFreeApp

PASS = []
FAIL = []


def check(name, cond, extra=""):
    (PASS if cond else FAIL).append(name)
    print(("PASS" if cond else "FAIL"), name, extra)


def pump(app, seconds=0.2):
    end = time.time() + seconds
    while time.time() < end:
        app.root.update()
        time.sleep(0.03)


def main():
    app = WordFreeApp()
    ed = app.editor
    t = ed.text
    pump(app, 0.5)
    check("app started", True)

    # --- ввод текста и проверка орфографии
    t.delete("1.0", "end")
    t.insert("1.0", "Привет мир. Это првиет тест.")
    t.edit_modified(True)
    ed._on_modified()
    pump(app, 0.1)
    ed._start_check()
    pump(app, 3.0)
    errs = t.tag_ranges("err_spell")
    check("spell error flagged", len(errs) >= 2, str(errs))
    cands = ed.spell.candidates("првиет")
    check("spell candidates", "привет" in cands, str(cands[:5]))

    # --- форматирование
    t.tag_add("sel", "1.0", "1.5")
    ed.toggle_char("bold")
    check("bold applied", bool(t.tag_ranges("b") and t.tag_ranges("b")[0] is not None))
    ed.toggle_char("bold")
    t.tag_remove("sel", "1.0", "end")

    # отмена/повтор
    ed.set_align("center")
    check("align center", "al_center" in t.tag_names("1.0"))
    ed.undo()
    check("undo align", "al_center" not in t.tag_names("1.0"))
    ed.redo()
    check("redo align", "al_center" in t.tag_names("1.0"))
    ed.undo()

    # --- списки
    t.tag_add("sel", "1.0", "1.end")
    ed.toggle_list("bullet")
    check("bullet list", "list_bullet" in t.tag_names("1.0"))
    check("bullet prefix", t.get("1.0", "1.6") in ("• Прив", "• При"), repr(t.get("1.0", "1.6")))
    ed.toggle_list("bullet")
    t.tag_remove("sel", "1.0", "end")

    # --- таблица
    tid = ed.insert_table(2, 3)
    check("table created", tid in ed.tables)
    ed.tables[tid]["cells"][0][0].insert(0, "Ячейка А1")
    blocks = ed.blocks()
    tbl_blocks = [b for b in blocks if b["type"] == "table"]
    check("table in blocks", len(tbl_blocks) == 1, str(len(tbl_blocks)))
    ed.table_add_row(tid)
    check("row added", ed.tables[tid]["rows"] == 3)

    # --- картинка
    from PIL import Image
    img_path = os.path.join(tempfile.mkdtemp(), "test.png")
    Image.new("RGB", (200, 100), (30, 120, 200)).save(img_path)
    t.mark_set("insert", "end-1c")
    ed.insert_image(img_path, 2.0)
    check("image inserted", len(ed.images) == 1)

    # --- сохранение docx
    docx_path = os.path.join(tempfile.mkdtemp(), "test.docx")
    t.mark_set("insert", "end-1c")
    ok = app._save_to(docx_path)
    check("save docx", ok and os.path.exists(docx_path), f"size={os.path.getsize(docx_path) if os.path.exists(docx_path) else 0}")
    from docx import Document
    d = Document(docx_path)
    texts = "\n".join(p.text for p in d.paragraphs)
    check("docx has text", "првиет" in texts, repr(texts[:80]))
    check("docx has table", len(d.tables) == 1, str(len(d.tables)))
    if d.tables:
        check("docx cell text", d.tables[0].cell(0, 0).text == "Ячейка А1",
              repr(d.tables[0].cell(0, 0).text))

    # --- PDF
    pdf_path = os.path.join(tempfile.mkdtemp(), "test.pdf")
    from wordfree import pdf_export
    margins = app.settings.get("margins_cm", [2, 2, 2, 2])
    pdf_export.export_pdf(t, ed.tables, ed.images, pdf_path,
                          title="Тест", margins_cm=margins)
    check("pdf exported", os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 1000,
          f"size={os.path.getsize(pdf_path) if os.path.exists(pdf_path) else 0}")

    # --- открытие docx
    app.open_path(docx_path)
    pump(app, 0.3)
    content = ed.text.get("1.0", "end-1c")
    check("open docx text", "првиет" in content, repr(content[:80]))
    check("open docx table", len(ed.tables) == 1, str(len(ed.tables)))

    # --- удаление таблицы (после переоткрытия — новый id)
    if ed.tables:
        tid2 = list(ed.tables.keys())[0]
        check("reopened cell text",
              ed.tables[tid2]["cells"][0][0].get() == "Ячейка А1",
              repr(ed.tables[tid2]["cells"][0][0].get()))
        ed.delete_table(tid2)
        check("table deleted", tid2 not in ed.tables)
    else:
        check("table deleted", False, "no table after open")

    # --- txt roundtrip
    txt_path = os.path.join(tempfile.mkdtemp(), "note.txt")
    ok = app._save_to(txt_path)
    check("save txt", ok)
    with open(txt_path, encoding="utf-8") as f:
        data = f.read()
    check("txt content", "првиет" in data)

    # --- поиск
    found = ed.find_next("мир")
    check("find next", found)

    # --- грамматика (правила точно; LT — если успел)
    issues_rules = __import__("wordfree.grammar", fromlist=["check_rules"]).check_rules(
        "это   тест  ,хороший.\n\n")
    check("grammar rules", len(issues_rules) >= 3, str(len(issues_rules)))

    # --- статистика
    st = ed.stats()
    check("stats words", st["words"] > 3, str(st))

    pump(app, 0.5)
    print()
    print(f"=== PASS {len(PASS)} / FAIL {len(FAIL)} ===")
    if FAIL:
        print("FAILED:", ", ".join(FAIL))
    app.root.destroy()
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
